"""Extract QA sign-off fields from Word documents (ported from backend 7)."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

from app.models.qa_signoff import SignOffRequest


def extract_docx_text(content: bytes) -> str:
    """Extract all meaningful text from a Word document."""
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
        lines: list[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_values = []
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        row_values.append(text)
                if row_values:
                    lines.append(" | ".join(row_values))

        return "\n".join(lines)
    except Exception as exc:
        raise ValueError(f"Could not read Word document: {exc}") from exc


def _known_field_labels() -> set[str]:
    return {
        "pr title",
        "test plan status",
        "open blocker or critical bugs",
        "pr tags",
        "summary",
        "expected result",
        "test environment",
        "qa validation",
        "test result",
        "defects identified",
        "release recommendation",
    }


def _is_field_label_only(line: str) -> bool:
    stripped = line.strip()
    if not stripped or ":" in stripped or "|" in stripped:
        return False
    return stripped.lower() in _known_field_labels()


def extract_field(text: str, field_name: str) -> str | None:
    """Extract a field value from inline, table, or label-on-next-line layouts."""
    lines = text.splitlines()
    inline_pattern = rf"^\s*{re.escape(field_name)}\s*:\s*(.+?)\s*$"
    pipe_pattern = rf"^\s*{re.escape(field_name)}\s*\|\s*(.+?)\s*$"
    label_pattern = rf"^\s*{re.escape(field_name)}\s*$"

    for index, line in enumerate(lines):
        match = re.match(inline_pattern, line, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip()

        match = re.match(pipe_pattern, line, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip()

        if re.match(label_pattern, line, flags=re.IGNORECASE):
            for next_line in lines[index + 1 :]:
                candidate = next_line.strip()
                if not candidate:
                    continue
                if _is_field_label_only(candidate):
                    break
                return candidate
    return None


def extract_pr_title(text: str) -> str:
    value = extract_field(text, "PR Title")
    if not value:
        raise ValueError("PR Title was not found in the Word document.")
    return value


def extract_test_plan_status(text: str) -> str:
    value = extract_field(text, "Test Plan Status")
    if not value:
        raise ValueError("Test Plan Status was not found in the Word document.")
    return value


def extract_open_bugs(text: str) -> bool:
    value = extract_field(text, "Open Blocker or Critical Bugs")
    if not value:
        raise ValueError("Open Blocker or Critical Bugs was not found in the Word document.")

    normalized = value.strip().lower()
    true_values = {"true", "yes", "y", "1", "open"}
    false_values = {"false", "no", "n", "0", "none", "closed"}

    if normalized in true_values:
        return True
    if normalized in false_values:
        return False

    raise ValueError(
        "Open Blocker or Critical Bugs must be True/False, Yes/No, or Open/Closed."
    )


def extract_tags(text: str) -> list[str]:
    value = extract_field(text, "PR Tags")
    if not value or value.lower() == "none":
        return []
    return [tag.strip() for tag in value.split(",") if tag.strip()]


def extract_summary(text: str) -> str:
    lines = text.splitlines()
    summary_started = False
    summary_lines: list[str] = []

    known_fields = {
        "PR Title",
        "Test Plan Status",
        "Open Blocker or Critical Bugs",
        "PR Tags",
        "Summary",
        "Expected Result",
        "Test Environment",
        "QA Validation",
        "Test Result",
        "Defects Identified",
        "Release Recommendation",
    }

    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue

        if re.match(r"^\s*Summary\s*:\s*(.*)$", stripped, flags=re.IGNORECASE):
            summary_started = True
            value = re.sub(r"^\s*Summary\s*:\s*", "", stripped, flags=re.IGNORECASE)
            if value:
                summary_lines.append(value)
            continue

        if re.match(r"^\s*Summary\s*$", stripped, flags=re.IGNORECASE):
            summary_started = True
            for next_line in lines[index + 1 :]:
                candidate = next_line.strip()
                if not candidate:
                    continue
                if _is_field_label_only(candidate):
                    break
                summary_lines.append(candidate)
            break

        if summary_started:
            is_known_field = any(
                re.match(rf"^\s*{re.escape(field)}\s*:", stripped, flags=re.IGNORECASE)
                or re.match(rf"^\s*{re.escape(field)}\s*$", stripped, flags=re.IGNORECASE)
                for field in known_fields
                if field != "Summary"
            )
            if is_known_field:
                break
            summary_lines.append(stripped)

    summary = " ".join(summary_lines).strip()
    return summary or "No summary provided."


def build_signoff_request_from_docx(text: str) -> SignOffRequest:
    return SignOffRequest(
        pr_title=extract_pr_title(text),
        test_plan_status=extract_test_plan_status(text),
        open_blocker_or_critical_bugs=extract_open_bugs(text),
        pr_tags=extract_tags(text),
        summary=extract_summary(text),
    )


def parse_signoff_attachment(content: bytes, filename: str) -> SignOffRequest:
    """Parse a QA sign-off attachment (.docx or .json) into SignOffRequest."""
    lower_name = filename.lower()

    if lower_name.endswith(".json"):
        try:
            payload = json.loads(content.decode("utf-8", errors="replace"))
        except json.JSONDecodeError as exc:
            raise ValueError(f"Uploaded JSON file must contain valid JSON. Error: {exc}") from exc
        return SignOffRequest.model_validate(payload)

    if lower_name.endswith(".docx"):
        document_text = extract_docx_text(content)
        if not document_text.strip():
            raise ValueError("The Word document does not contain readable text.")
        return build_signoff_request_from_docx(document_text)

    raise ValueError("QA sign-off attachment must be a .docx or .json file.")


def read_signoff_attachment(path: Path, filename: str) -> SignOffRequest:
    if not path.is_file():
        raise ValueError("QA sign-off attachment file was not found on disk.")
    return parse_signoff_attachment(path.read_bytes(), filename)
